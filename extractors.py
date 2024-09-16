import os
import PyPDF2
import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import openpyxl
from PIL import Image
import pytesseract
import io
import unicodedata
from concurrent.futures import ThreadPoolExecutor, TimeoutError
import ebooklib
from ebooklib import epub
from bs4 import BeautifulSoup
import mobi
from logging_setup import logger
import speech_recognition as sr
from pydub import AudioSegment
import json
from vosk import Model, KaldiRecognizer, SetLogLevel, BatchRecognizer, BatchModel, GpuInit
import wave
from langdetect import detect
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading
import numpy as np
import time

def extract_text(file_path: str) -> str:
    _, ext = os.path.splitext(file_path.lower())
    extraction_functions = {
        '.pdf': extract_pdf,
        '.docx': extract_word,
        '.doc': extract_word,
        '.xlsx': extract_excel,
        '.xls': extract_excel,
        '.png': extract_image,
        '.jpg': extract_image,
        '.jpeg': extract_image,
        '.mobi': "skip", # DOESN'T WORK YET
        '.epub': extract_epub,
        '.mp3': extract_audio_text,
        '.wav': extract_audio_text,
        '.ogg': extract_audio_text,
        '.flac': extract_audio_text,
    }
    extract_func = extraction_functions.get(ext)
    if extract_func == "skip":
        return ''
    
    if extract_func:
        with ThreadPoolExecutor() as executor:
            future = executor.submit(extract_func, file_path)
            try:
                text = future.result(timeout=10.0)  # 10 seconds timeout
                return unicodedata.normalize('NFC', text)
            except TimeoutError:
                logger.error(f"Timeout extracting text from {file_path}")
            except Exception as e:
                logger.error(f"Error extracting text from {file_path}: {str(e)}")
    else:
        logger.warning(f"Unsupported file type: {ext} for file {file_path}")
    return ''

def extract_pdf(file_path):
    text = ""
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        for i, page in enumerate(reader.pages):
            if i > 250:
                break
            text += page.extract_text() + "\n"
    return text

def extract_word(file_path):
    doc = docx.Document(file_path)
    text = " ".join([paragraph.text for paragraph in doc.paragraphs])
    image_text = ""
    for rel in doc.part.rels.values():
        if rel.reltype == RT.IMAGE:
            image_stream = rel.target_part.blob
            image = Image.open(io.BytesIO(image_stream))
            image_text += pytesseract.image_to_string(image) + " "
    return text + " " + image_text

def extract_excel(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = ""
    for sheet in wb:
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i > 1000:
                break
            text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text

def extract_image(file_path):
    return pytesseract.image_to_string(Image.open(file_path))

def extract_mobi(file_path):
    try:
        extraction_result = mobi.extract(file_path)
        
        if isinstance(extraction_result, tuple):
            extracted_path = extraction_result[1]
        else:
            extracted_path = extraction_result

        text_buffer = io.StringIO()
        
        for root, dirs, files in os.walk(extracted_path):
            for file in files:
                if file.endswith('.html') or file.endswith('.htm'):
                    with open(os.path.join(root, file), 'r', encoding='utf-8', errors='ignore') as f:
                        soup = BeautifulSoup(f, 'html.parser')
                        text_buffer.write(soup.get_text() + "\n")
        
        text = text_buffer.getvalue()
        logger.debug(f"Extracted text from MOBI: {text[:100]}...")
        return text
    except Exception as e:
        logger.error(f"Error extracting MOBI file {file_path}: {str(e)}")
        return ""
    finally:
        # Clean up extracted files if they exist
        if 'extracted_path' in locals():
            for root, dirs, files in os.walk(extracted_path, topdown=False):
                for file in files:
                    os.remove(os.path.join(root, file))
                for dir in dirs:
                    os.rmdir(os.path.join(root, dir))
            os.rmdir(extracted_path)

def extract_epub(file_path):
    try:
        book = epub.read_epub(file_path)
        text = ""
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), 'html.parser')
                text += soup.get_text() + "\n"
        logger.debug(f"Extracted text from EPUB: {text[:100]}...")
        return text
    except Exception as e:
        logger.error(f"Error extracting EPUB file {file_path}: {str(e)}")
        return ""
    
# Global variable to store the Vosk model
vosk_model = None
model_lock = threading.Lock()
audio_files_queue = []

def load_vosk_model(model_path="vosk-model-small-en-us-0.15"):
    global vosk_model
    with model_lock:
        if vosk_model is None:
            if not os.path.exists(model_path):
                logger.error(f"Vosk model not found. Please download it from https://alphacephei.com/vosk/models and extract to {model_path}")
                return None
            logger.info(f"Loading Vosk model from {model_path}")
            vosk_model = Model(model_path)
            logger.info("Vosk model loaded successfully")
    return vosk_model

def process_audio_chunk(chunk, sample_rate):
    model = load_vosk_model()
    if model is None:
        return "Vosk model not loaded"

    rec = KaldiRecognizer(model, sample_rate)
    rec.SetWords(True)

    if rec.AcceptWaveform(chunk):
        result = json.loads(rec.Result())
    else:
        result = json.loads(rec.FinalResult())

    return result.get('text', '')

def extract_audio_text(file_path):
    global audio_files_queue
    filename = os.path.basename(file_path)
    print(f"Extracting text from audio file: {filename}")
    try:
        lang = detect(filename)
    except:
        lang = "unknown"
        
    if lang != 'en':
        logger.warning(f"Unsupported language: {lang} for audio file {filename}. Skipping.")
        return " "
    
    audio_files_queue.append(file_path)
    return ""

def process_audio_file(file_path, chunk_duration_ms=60000, max_duration=30, max_workers=3, sample_rate=16000):
    _, ext = os.path.splitext(file_path.lower())
    
    logger.info(f"Processing audio file: {file_path}")

    try:
        # Load audio file
        audio = AudioSegment.from_file(file_path, format=ext[1:])
        logger.info(f"Audio file loaded. Duration: {len(audio)/1000:.2f} seconds")

        # Downsample and convert to mono
        audio = audio.set_frame_rate(sample_rate).set_channels(1)
        
        audio_duration_minutes = len(audio)/60000        
        
        if audio_duration_minutes <= max_duration:
            skip = 1
        else:
            skip = int(audio_duration_minutes // max_duration)
            
        logger.info(f"Audio duration: {audio_duration_minutes:.2f} minutes. Skip: {skip}")
        
        # Split audio into chunks
        chunks = [audio[i:i+chunk_duration_ms] for i in range(0, len(audio), chunk_duration_ms)]
        logger.info(f"Audio split into {len(chunks)} chunks of {chunk_duration_ms}ms each")

        # Process chunks concurrently
        results = []
        start = time.time()
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = []
            for i, chunk in enumerate(chunks):
                if i % skip != 0:
                    continue
                chunk_data = np.frombuffer(chunk.raw_data, dtype=np.int16)
                futures.append(executor.submit(process_audio_chunk, chunk_data.tobytes(), sample_rate))
            
            for i, future in enumerate(as_completed(futures)):
                try:
                    result = future.result()
                    results.append(result)
                    time_remaining = (time.time() - start) / (i + 1) * (len(futures) - i - 1)
                    logger.info(f"Processed chunk {i+1}/{len(futures)}. Remaining time: {time_remaining:.2f} seconds")
                except Exception as e:
                    logger.error(f"Error processing chunk {i+1}/{len(futures)}: {str(e)}")

        # Combine results
        text = ' '.join(results)
        logger.info(f"Audio processing completed. Extracted text length: {len(text)} characters")

        return text

    except Exception as e:
        logger.error(f"Error processing audio file {file_path}: {str(e)}")
        return ""

def process_audio_queue(max_workers=3):
    global audio_files_queue
    
    if not audio_files_queue:
        logger.info("No audio files in the queue to process.")
        return

    logger.info(f"Processing {len(audio_files_queue)} audio files from the queue.")
    
    results = {}
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_to_file = {executor.submit(process_audio_file, file_path): file_path for file_path in audio_files_queue}
        for future in as_completed(future_to_file):
            file_path = future_to_file[future]
            try:
                text = future.result()
                results[file_path] = text
                logger.info(f"Processed {file_path}: {text[:100]}...")  # Print first 100 characters of each result
            except Exception as e:
                logger.error(f"Error processing {file_path}: {str(e)}")
    
    audio_files_queue.clear()
    return results

def sanitize_filename(filename):
    return ''.join(c for c in filename if c.isalnum() or c in ('-', '_')).rstrip()

def process_files(directory):
    for root, _, files in os.walk(directory):
        for file in files:
            file_path = os.path.join(root, file)
            logger.info(f"Indexing: {file_path}")
        
            text = extract_text(file_path)
            
            if not text:
                logger.warning(f"No content extracted from: {file_path}")
            else:
                output_filename = sanitize_filename(os.path.splitext(file)[0]) + ".txt"
                output_path = os.path.join(root, output_filename)
                with open(output_path, 'w', encoding='utf-8') as f:
                    f.write(text)
                logger.info(f"Extracted text saved to: {output_path}")
                
if __name__ == '__main__':
    text = extract_audio_text('/Users/blaz/Library/CloudStorage/OneDrive-Personal/Dokumenti/[01] Imported/Avdioknjige/Getting Things Done.mp3')
    print(text)